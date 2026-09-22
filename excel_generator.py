import base64
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from utils import formatar_moeda


def gerar_excel_cotacao(dados_gerais, df_resultados):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comparativo de Frete"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    sub_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    best_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    info_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    font_title = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_normal = Font(name="Segoe UI", size=10, color="1E293B")
    font_best = Font(name="Segoe UI", size=10, bold=True, color="166534")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = "SISTEMA DE COTAÇÕES LOGÍSTICAS - NEXT CABLE"
    title_cell.font = font_title
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:G2")
    ws["A2"].value = "PARÂMETROS DA COTAÇÃO"
    ws["A2"].font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    ws["A2"].fill = info_fill
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    ws["A3"] = "Data/Hora:"
    ws["B3"] = dados_gerais.get("data_hora", "")
    ws["A4"] = "Origem:"
    ws["B4"] = f"{dados_gerais.get('origem_cep', '')} ({dados_gerais.get('origem_texto', '')})"
    ws["A5"] = "Destino:"
    ws["B5"] = f"{dados_gerais.get('destino_cep', '')} ({dados_gerais.get('destino_texto', '')})"

    ws["D3"] = "Valor da NF:"
    ws["E3"] = f"R$ {dados_gerais.get('valor_nf', 0):,.2f}"
    ws["D4"] = "Peso Real:"
    ws["E4"] = f"{dados_gerais.get('peso', 0):.2f} kg"
    ws["D5"] = "Peso Tarifado:"
    ws["E5"] = f"{dados_gerais.get('peso_tarifado', 0):.2f} kg (Cubado: {dados_gerais.get('peso_cubado', 0):.2f} kg)"

    for r in range(3, 6):
        ws[f"A{r}"].font = font_bold
        ws[f"B{r}"].font = font_normal
        ws[f"D{r}"].font = font_bold
        ws[f"E{r}"].font = font_normal

    headers = ["Classificação", "Transportadora", "Nº Cotação", "Tipo", "Valor Frete", "Prazo", "Status / Detalhe"]
    start_row = 7
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = sub_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin
    ws.row_dimensions[start_row].height = 24

    for idx, row in df_resultados.iterrows():
        curr_row = start_row + 1 + idx
        is_best = (idx == 0 and pd.notna(row.get("Valor Frete (R$)")))

        pos_txt = "★ MELHOR" if is_best else (f"{idx+1}º Lugar" if pd.notna(row.get("Valor Frete (R$)")) else "—")
        val_txt = f"R$ {row['Valor Frete (R$)']:,.2f}" if pd.notna(row.get("Valor Frete (R$)")) else "—"
        prazo_num = int(row.get("Prazo (Dias Úteis)") or 0) if pd.notna(row.get("Prazo (Dias Úteis)")) else 0
        prazo_txt = f"{prazo_num} dias úteis" if prazo_num > 0 else "A confirmar"
        tipo_txt = "API Oficial" if row.get("Fonte") == "api" and pd.notna(row.get("Valor Frete (R$)")) else (
            "Tabela Estimada" if row.get("Fonte") == "tabela" and pd.notna(row.get("Valor Frete (R$)")) else "—"
        )

        cells = [
            ws.cell(row=curr_row, column=1, value=pos_txt),
            ws.cell(row=curr_row, column=2, value=str(row.get("Transportadora", ""))),
            ws.cell(row=curr_row, column=3, value=str(row.get("Nº Cotação", ""))),
            ws.cell(row=curr_row, column=4, value=tipo_txt),
            ws.cell(row=curr_row, column=5, value=val_txt),
            ws.cell(row=curr_row, column=6, value=prazo_txt),
            ws.cell(row=curr_row, column=7, value=str(row.get("Status", ""))),
        ]

        for cell in cells:
            cell.font = font_best if is_best else font_normal
            if is_best:
                cell.fill = best_fill
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

        cells[0].alignment = Alignment(horizontal="center", vertical="center")
        cells[4].alignment = Alignment(horizontal="right", vertical="center")
        cells[5].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[curr_row].height = 20

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(start_row, start_row + 1 + len(df_resultados)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def gerar_excel_entregas(df_entregas):
    """Gera planilha profissional de monitoramento de entregas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monitoramento de Entregas"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    font_title = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_normal = Font(name="Segoe UI", size=9, color="1E293B")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"), right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"), bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = f"RELATÓRIO DE MONITORAMENTO DE ENTREGAS - NEXT CABLE ({datetime.now():%d/%m/%Y %H:%M})"
    t.font = font_title
    t.fill = header_fill
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    colunas = ["CNPJ", "Nota Fiscal", "Status", "Previsão de Entrega", "Data de Entrega", "Última Ocorrência", "Data Ocorrência", "Destino", "Valor Frete"]
    for col_idx, col_name in enumerate(colunas, 1):
        c = ws.cell(row=2, column=col_idx, value=col_name)
        c.font = font_header
        c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_thin
    ws.row_dimensions[2].height = 22

    fill_entregue = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_transito = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_erro = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    for row_idx, r in df_entregas.iterrows():
        curr_row = 3 + row_idx
        st_text = str(r.get("Status", ""))
        row_fill = None
        if "entregue" in st_text.lower():
            row_fill = fill_entregue
        elif "trânsito" in st_text.lower() or "transito" in st_text.lower() or "saiu" in st_text.lower():
            row_fill = fill_transito
        elif "ocorrência" in st_text.lower() or "erro" in st_text.lower() or "retido" in st_text.lower():
            row_fill = fill_erro

        for col_idx, col_name in enumerate(colunas, 1):
            val = str(r.get(col_name, "") if pd.notna(r.get(col_name)) else "")
            c = ws.cell(row=curr_row, column=col_idx, value=val)
            c.font = font_normal
            c.border = border_thin
            c.alignment = Alignment(vertical="center")
            if row_fill and col_name == "Status":
                c.fill = row_fill
        ws.row_dimensions[curr_row].height = 19

    for col_idx in range(1, len(colunas) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, 3 + len(df_entregas)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

